#!/usr/bin/env python3
import os
import re
import sys
import zipfile
import subprocess

try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except:
    pdf_extract_text = None

try:
    from docx import Document
except:
    Document = None

try:
    import openpyxl
except:
    openpyxl = None

try:
    import xlrd
except:
    xlrd = None

try:
    from odf.opendocument import load as odf_load
    from odf.text import P as odf_P
    from odf.table import TableCell
except:
    odf_load = None


EMAIL_REGEX = re.compile(
    r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}",
    re.IGNORECASE
)

TEXT_EXT = {".txt", ".csv", ".tsv", ".log", ".json", ".xml",
            ".html", ".htm", ".md", ".yaml", ".yml", ".ini", ".rtf"}

PDF_EXT = {".pdf"}
DOC_EXT = {".doc"}
DOCX_EXT = {".docx"}
XLSX_EXT = {".xlsx"}
XLS_EXT = {".xls"}
ODF_EXT = {".odt", ".ods", ".odp", ".odg", ".odf"}

ZIP_EXT = {".zip"}


def extract_emails(text):
    return set(EMAIL_REGEX.findall(text))


def read_pdf(path):
    if not pdf_extract_text:
        return set()
    try:
        return extract_emails(pdf_extract_text(path))
    except:
        return set()


def read_docx(path):
    if not Document:
        return set()
    try:
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        return extract_emails(text)
    except:
        return set()


def read_doc(path):
    """Use macOS textutil (native) to convert .doc → .txt"""
    try:
        out = subprocess.check_output(
            ["textutil", "-convert", "txt", path, "-stdout"],
            stderr=subprocess.DEVNULL
        )
        return extract_emails(out.decode("utf-8", errors="ignore"))
    except:
        return set()


def read_xlsx(path):
    if not openpyxl:
        return set()
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        emails = set()
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        emails |= extract_emails(cell)
        return emails
    except:
        return set()


def read_xls(path):
    if not xlrd:
        return set()
    try:
        wb = xlrd.open_workbook(path)
        emails = set()
        for sheet in wb.sheets():
            for rowx in range(sheet.nrows):
                for cell in sheet.row_values(rowx):
                    if isinstance(cell, str):
                        emails |= extract_emails(cell)
        return emails
    except:
        return set()


def read_odf(path):
    if not odf_load:
        return set()
    try:
        doc = odf_load(path)
        parts = []
        for p in doc.getElementsByType(odf_P):
            parts.append(str(p))
        for cell in doc.getElementsByType(TableCell):
            txt = "".join(
                n.data for n in getattr(cell, "childNodes", [])
                if hasattr(n, "data")
            )
            parts.append(txt)
        return extract_emails("\n".join(parts))
    except:
        return set()


def read_zip(path):
    emails = set()
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if name.lower().endswith(tuple(TEXT_EXT)):
                    try:
                        content = z.read(name).decode("utf-8", errors="ignore")
                        emails |= extract_emails(content)
                    except:
                        pass
        return emails
    except:
        return set()


def read_text(path):
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return extract_emails(f.read())
    except:
        return set()


def process_file(path):
    ext = os.path.splitext(path)[1].lower()

    if ext in TEXT_EXT:
        return read_text(path)
    if ext in PDF_EXT:
        return read_pdf(path)
    if ext in DOCX_EXT:
        return read_docx(path)
    if ext in DOC_EXT:
        return read_doc(path)
    if ext in XLSX_EXT:
        return read_xlsx(path)
    if ext in XLS_EXT:
        return read_xls(path)
    if ext in ODF_EXT:
        return read_odf(path)
    if ext in ZIP_EXT:
        return read_zip(path)

    return set()


def scan(root):
    for dirpath, _, files in os.walk(root):
        for f in files:
            path = os.path.join(dirpath, f)
            emails = process_file(path)

            if len(emails) >= 20:
                print(f"\n=== {path} ===")
                print(f"{len(emails)} emails trouvés :")
#                 for e in sorted(emails):
#                     print("  -", e)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python3 list-file-with-emails.py <répertoire>")
        sys.exit(1)

    scan(sys.argv[1])